from __future__ import annotations

import os
import re
from collections.abc import Iterator
from pathlib import Path
from typing import TYPE_CHECKING

from core import config
from core.errors import InvalidContent, UnsupportedMedia

if TYPE_CHECKING:
    pass

_WORD = re.compile(r"\S+")
_TEXT_BLOCK_CHARS = 1 << 20
_MAX_TOKEN_CHARS = 1 << 16
_PDF_MIN_TEXT_CHARS = 40
_OCR_DPI = 200


def detect_text_encoding(path: str | os.PathLike) -> str:
    from charset_normalizer import from_path

    match = from_path(path).best()
    if match is None:
        raise InvalidContent("could not detect a text encoding")
    encoding = match.encoding
    if not encoding:
        raise InvalidContent("could not detect a text encoding")
    return encoding


def _read_head(path: str | os.PathLike, n: int = 16) -> bytes:
    with open(path, "rb") as f:
        return f.read(n)


def _require_pdf(path: str | os.PathLike) -> None:
    if _read_head(path, 5) != b"%PDF-":
        raise UnsupportedMedia("missing %PDF- header")


def _require_zip(path: str | os.PathLike, label: str) -> None:
    if _read_head(path, 4) != b"PK\x03\x04":
        raise UnsupportedMedia(f"not a valid {label} archive")


def _open_image(path: str | os.PathLike):
    from PIL import Image

    try:
        im = Image.open(path)
        im.load()
        return im
    except OSError as e:
        raise UnsupportedMedia(f"unreadable image: {e}") from e


def _ocr_image(im) -> str:
    import pytesseract

    try:
        return pytesseract.image_to_string(im) or ""
    except pytesseract.TesseractNotFoundError as e:
        raise InvalidContent(
            "OCR requires Tesseract (install tesseract-ocr on the system PATH)"
        ) from e
    except OSError as e:
        raise InvalidContent(f"OCR failed: {e}") from e


def _words_from_text(text: str, page: int | None) -> Iterator[tuple[str, int | None]]:
    for m in _WORD.finditer(text):
        yield m.group(), page


def inspect_file(path: str | os.PathLike, ext: str) -> str | None:
    ext = ext.lower()
    if ext == ".pdf":
        _require_pdf(path)
        return None
    if ext in config.IMAGE_EXT:
        with _open_image(path):
            pass
        return None
    if ext in config.ZIP_DOC_EXT:
        labels = {
            ".docx": "Word",
            ".pptx": "PowerPoint",
            ".xlsx": "Excel",
            ".odt": "OpenDocument",
            ".epub": "EPUB",
        }
        _require_zip(path, labels.get(ext, "document"))
        return None
    if ext in config.TEXT_EXT or ext in {".rtf", ".html", ".htm"}:
        return detect_text_encoding(path)
    raise UnsupportedMedia(f"no inspector for extension {ext}")


def iter_text_words(path: str | os.PathLike, encoding: str) -> Iterator[tuple[str, int | None]]:
    tail = ""
    with open(path, encoding=encoding, errors="replace") as f:
        while block := f.read(_TEXT_BLOCK_CHARS):
            block = tail + block
            tail = ""
            for m in _WORD.finditer(block):
                if m.end() == len(block) and len(block) - m.start() < _MAX_TOKEN_CHARS:
                    tail = m.group()
                else:
                    yield m.group(), None
    if tail:
        yield tail, None


def iter_image_words(path: str | os.PathLike) -> Iterator[tuple[str, int | None]]:
    with _open_image(path) as im:
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        text = _ocr_image(im)
    yield from _words_from_text(text, None)


def iter_pdf_words(path: str | os.PathLike) -> Iterator[tuple[str, int | None]]:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(path)
        pages: list[tuple[str, int]] = []
        for page_no, page in enumerate(reader.pages, start=1):
            pages.append((page.extract_text() or "", page_no))
    except (PyPdfError, ValueError, KeyError, TypeError) as e:
        raise InvalidContent(f"unreadable PDF: {type(e).__name__}") from e

    joined = "".join(t for t, _ in pages)
    if len(joined.strip()) >= _PDF_MIN_TEXT_CHARS:
        for text, page_no in pages:
            yield from _words_from_text(text, page_no)
        return
    yield from _iter_pdf_ocr_words(path)


def _iter_pdf_ocr_words(path: str | os.PathLike) -> Iterator[tuple[str, int | None]]:
    import fitz

    try:
        doc = fitz.open(path)
    except Exception as e:
        raise InvalidContent(f"PDF OCR failed to open file: {type(e).__name__}") from e
    try:
        scale = _OCR_DPI / 72.0
        matrix = fitz.Matrix(scale, scale)
        for page_no, page in enumerate(doc, start=1):
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            from PIL import Image

            im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            text = _ocr_image(im)
            yield from _words_from_text(text, page_no)
    finally:
        doc.close()


def iter_docx_words(path: str | os.PathLike) -> Iterator[tuple[str, int | None]]:
    from docx import Document

    try:
        doc = Document(path)
    except Exception as e:
        raise InvalidContent(f"unreadable Word document: {type(e).__name__}") from e
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    yield from _words_from_text("\n".join(parts), None)


def iter_odt_words(path: str | os.PathLike) -> Iterator[tuple[str, int | None]]:
    from odf import teletype
    from odf.opendocument import load

    try:
        doc = load(path)
    except Exception as e:
        raise InvalidContent(f"unreadable ODT: {type(e).__name__}") from e
    yield from _words_from_text(teletype.extractText(doc), None)


def iter_pptx_words(path: str | os.PathLike) -> Iterator[tuple[str, int | None]]:
    from pptx import Presentation

    try:
        prs = Presentation(path)
    except Exception as e:
        raise InvalidContent(f"unreadable PowerPoint: {type(e).__name__}") from e
    for slide_no, slide in enumerate(prs.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                chunks.append(shape.text)
        yield from _words_from_text("\n".join(chunks), slide_no)


def iter_xlsx_words(path: str | os.PathLike) -> Iterator[tuple[str, int | None]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise InvalidContent(f"unreadable Excel workbook: {type(e).__name__}") from e
    try:
        for sheet_no, sheet in enumerate(wb.worksheets, start=1):
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                for val in row:
                    if val is not None and str(val).strip():
                        rows.append(str(val))
            yield from _words_from_text(" ".join(rows), sheet_no)
    finally:
        wb.close()


def iter_rtf_words(path: str | os.PathLike, encoding: str) -> Iterator[tuple[str, int | None]]:
    from striprtf.striprtf import rtf_to_text

    try:
        raw = Path(path).read_text(encoding=encoding, errors="replace")
        text = rtf_to_text(raw)
    except Exception as e:
        raise InvalidContent(f"unreadable RTF: {type(e).__name__}") from e
    yield from _words_from_text(text, None)


def iter_html_words(path: str | os.PathLike, encoding: str) -> Iterator[tuple[str, int | None]]:
    from bs4 import BeautifulSoup

    try:
        html = Path(path).read_text(encoding=encoding, errors="replace")
        text = BeautifulSoup(html, "html.parser").get_text(separator=" ")
    except Exception as e:
        raise InvalidContent(f"unreadable HTML: {type(e).__name__}") from e
    yield from _words_from_text(text, None)


def iter_epub_words(path: str | os.PathLike) -> Iterator[tuple[str, int | None]]:
    from bs4 import BeautifulSoup
    from ebooklib import ITEM_DOCUMENT, epub

    try:
        book = epub.read_epub(str(path))
    except Exception as e:
        raise InvalidContent(f"unreadable EPUB: {type(e).__name__}") from e
    for item in book.get_items():
        if item.get_type() != ITEM_DOCUMENT:
            continue
        text = BeautifulSoup(item.get_content(), "html.parser").get_text(separator=" ")
        yield from _words_from_text(text, None)


def iter_document_words(
    path: str | os.PathLike, ext: str, encoding: str | None = None
) -> Iterator[tuple[str, int | None]]:
    ext = ext.lower()
    if ext == ".pdf":
        return iter_pdf_words(path)
    if ext in config.IMAGE_EXT:
        return iter_image_words(path)
    if ext == ".docx":
        return iter_docx_words(path)
    if ext == ".odt":
        return iter_odt_words(path)
    if ext == ".pptx":
        return iter_pptx_words(path)
    if ext == ".xlsx":
        return iter_xlsx_words(path)
    if ext == ".epub":
        return iter_epub_words(path)
    if ext == ".rtf":
        return iter_rtf_words(path, encoding or detect_text_encoding(path))
    if ext in {".html", ".htm"}:
        return iter_html_words(path, encoding or detect_text_encoding(path))
    if ext in config.TEXT_EXT:
        return iter_text_words(path, encoding or "utf-8")
    raise UnsupportedMedia(f"no extractor for extension {ext}")
